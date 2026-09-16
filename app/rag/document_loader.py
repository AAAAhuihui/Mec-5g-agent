from __future__ import annotations

import re
from pathlib import Path
from typing import Any


SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".pdf", ".docx"}


def resolve_source_dir(source_dir: str) -> Path:
    path = Path(source_dir)
    if not path.is_absolute():
        path = Path(__file__).resolve().parents[2] / path
    return path


def load_documents(source_dir: str) -> list[dict[str, Any]]:
    """Load documents as page-aware units."""
    root = resolve_source_dir(source_dir)
    if not root.exists():
        raise FileNotFoundError(f"source_dir not found: {root}")

    documents: list[dict[str, Any]] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        if path.suffix.lower() == ".pdf":
            documents.extend(_load_pdf_pages(path))
            continue
        text = load_document_text(path)
        if text.strip():
            documents.append(
                {
                    "content": text,
                    "source": path.name,
                    "path": str(path),
                    "metadata": {"document_type": path.suffix.lower().lstrip(".")},
                }
            )
    return documents


def load_document_text(path: Path) -> str:
    """Load a whole file as text (kept for legacy-label migration)."""
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        return _load_pdf_text_legacy(path)
    if suffix == ".docx":
        return _load_docx(path)
    raise ValueError(f"unsupported document type: {path.suffix}")


def _load_pdf_pages(path: Path) -> list[dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError("Reading PDF files requires: pip install pypdf") from exc

    reader = PdfReader(str(path))
    pages: list[dict[str, Any]] = []
    for index, page in enumerate(reader.pages, start=1):
        text = _clean_pdf_page(page.extract_text() or "", index)
        if not text:
            continue
        pages.append(
            {
                "content": text,
                "source": path.name,
                "path": str(path),
                "metadata": {"document_type": "pdf", "page": index},
            }
        )
    return pages


def _load_pdf_text_legacy(path: Path) -> str:
    """Reproduce the pre-v2 PDF text stream for migrating old chunk labels."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError("Reading PDF files requires: pip install pypdf") from exc

    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"# Page {index}\n{text.strip()}")
    return "\n\n".join(pages)


def _clean_pdf_page(text: str, page_number: int) -> str:
    """Remove repeated 3GPP page furniture without touching body text."""
    lines = [line.rstrip() for line in text.replace("\u00a0", " ").splitlines()]
    cleaned: list[str] = []
    header_pattern = re.compile(
        rf"^3GPP TS \d+\.\d+ V[^ ]+ \(\d{{4}}-\d{{2}}\)\s+{page_number}\s+Release \d+$",
        re.IGNORECASE,
    )
    for line in lines:
        stripped = line.strip()
        if stripped == "3GPP" or header_pattern.match(stripped):
            continue
        cleaned.append(stripped)

    output: list[str] = []
    previous_blank = False
    for line in cleaned:
        blank = not line
        if blank and previous_blank:
            continue
        output.append(line)
        previous_blank = blank
    return "\n".join(output).strip()


def _load_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("Reading docx files requires: pip install python-docx") from exc

    document = Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
